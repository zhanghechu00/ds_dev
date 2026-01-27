from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    # Create a blank presentation
    prs = Presentation()

    # Helper function to add a title slide
    def add_title_slide(title_text, subtitle_text):
        slide_layout = prs.slide_layouts[0] # 0 is Title Slide
        slide = prs.slides.add_slide(slide_layout)
        
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = title_text
        subtitle.text = subtitle_text

    # Helper function to add a content slide (Title + Bullets)
    def add_content_slide(title_text, content_list):
        slide_layout = prs.slide_layouts[1] # 1 is Title and Content
        slide = prs.slides.add_slide(slide_layout)
        
        # Set Title
        title = slide.shapes.title
        title.text = title_text
        
        # Set Content
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        tf.word_wrap = True
        
        for i, item in enumerate(content_list):
            if isinstance(item, tuple):
                text = item[0]
                level = item[1]
            else:
                text = item
                level = 0
            
            # Add paragraph (first one exists by default, others need adding)
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            
            p.text = text
            p.level = level
            p.font.size = Pt(24 - (level * 4)) # Resize based on level
            
            # Add some spacing
            p.space_after = Pt(10)

    # --- Slide 1: Cover ---
    add_title_slide(
        "企业级 AI Agent 智能体平台建设方案",
        "构建“数字员工”生产力引擎，驱动业务智能化转型\n\n汇报人：[您的名字/团队名称]\n日期：2026年1月"
    )

    # --- Slide 2: Executive Summary ---
    add_content_slide("核心愿景 (Executive Summary)", [
        "目标：打造企业统一的 AI 智能体开发与运行平台",
        "价值主张：",
        ("低门槛：让业务人员也能“雇佣”数字员工", 1),
        ("破孤岛：打通文档、数据库与遗留系统", 1),
        ("高效率：实现从“对话”到“行动”的自动化闭环", 1),
        "关键产出：",
        ("Agent 开发工坊 (可视编排)", 1),
        ("工具连接网关 (MCP Gateway)", 1),
        ("知识中台 (Knowledge Hub)", 1)
    ])

    # --- Slide 3: Why Now ---
    add_content_slide("现状与痛点 (Why Now?)", [
        "行业趋势：大模型从 Chat (闲聊) 进化为 Agent (行动)",
        "当前企业痛点：",
        ("能力割裂：文档知识与业务系统无法互通", 1),
        ("  (例：LLM 无法理解 C++ 处理的油气数据)", 2),
        ("僵化：传统 RPA 无法处理模糊指令", 1),
        ("  (例：“分析这份报告，如果没有结论则忽略”)", 2),
        ("管控缺失：缺乏统一的安全审计与权限管理", 1)
    ])

    # --- Slide 4: Architecture ---
    add_content_slide("平台总体架构 (High-Level Architecture)", [
        "[在此处插入：分层架构图]",
        "应用层：",
        ("智能客服、自动化办公、专业数据分析助手", 1),
        "功能层：",
        ("Agent 开发工坊、应用市场、管理后台", 1),
        "引擎层 (核心)：",
        ("编排引擎：任务规划与拆解 (ReAct)", 1),
        ("记忆系统：长短期记忆 (Vector DB)", 1),
        ("工具网关：MCP 标准连接器", 1),
        "模型层：GPT-4 / Claude / DeepSeek (私有化)"
    ])

    # --- Slide 5: Tool Gateway (MCP) ---
    add_content_slide("核心引擎亮点：通用工具网关", [
        "痛点解决：如何让 AI 调用本地复杂的 Python/C++ 业务代码？",
        "技术方案：基于 MCP (Model Context Protocol) 协议",
        "案例演示：",
        ("AI 智能体 -> 编排引擎", 1),
        ("-> 直接调度 backend.py (Python)", 1),
        ("-> 调用 ExportPetroFileV32.exe (遗留 C++ 模块)", 1),
        ("优势：解耦大模型与底层业务逻辑，支持异构语言环境", 0)
    ])

    # --- Slide 6: Knowledge Engine (RAG) ---
    add_content_slide("知识中台：RAG 检索增强", [
        "功能：让模型拥有企业私有知识",
        "流程：文档上传 -> 自动切片 -> 向量化存储 -> 语义检索 -> 回答生成",
        "特色能力：",
        ("支持多格式解析 (PDF, MD, Word, 代码)", 1),
        ("源头溯源：回答中包含原文引用链接，确保可信度", 1),
        ("知识应用：如“基于技术文档自动生成 PPT”", 1)
    ])

    # --- Slide 7: Tech Stack ---
    add_content_slide("技术栈选型 (Tech Stack)", [
        "后端：Python (FastAPI) + LangChain / LangGraph",
        "前端：React + 流程编排组件",
        "向量数据库：PGVector (推荐) 或 Milvus",
        "核心协议：",
        ("全面适配 MCP (Model Context Protocol)", 1),
        ("确保工具生态兼容性，方便未来扩展", 1),
        "部署架构：Docker 容器化 + K8s 集群"
    ])

    # --- Slide 8: Hardware Resources ---
    add_content_slide("基础设施配置 (Resource Requirements)", [
        "方案 A：混合云/API 模式 (推荐 - 高性价比)",
        ("应用服务器 (2台)：运行 Agent 逻辑、MCP Server", 1),
        ("向量库服务器 (1台)：大内存，存储企业知识库", 1),
        ("遗留系统节点 (1台 Windows)：运行旧版 C++ 工具", 1),
        "方案 B：私有化部署 (数据严控)",
        ("LLM 推理集群：需配备 NVIDIA A100/H800 或 RTX 4090 集群", 1),
        ("显存需求：支持 70B+ 参数模型 (如 DeepSeek-70B)", 1)
    ])

    # --- Slide 9: Budget ---
    add_content_slide("成本构成与预算 (Budget Breakdown)", [
        "一次性投入 (CAPEX)：",
        ("硬件：服务器与 GPU 采购费用", 1),
        ("软件：OS、数据库授权、数据清洗人力", 1),
        ("研发人力：架构师、前后端、Prompt 工程师", 1),
        "持续运营 (OPEX)：",
        ("Token 费用：公有云 API 调用费 (按量)", 1),
        ("能耗与维保：私有化机房电费、系统运维", 1),
        ("持续迭代：Agent 技能更新、知识库维护", 1)
    ])

    # --- Slide 10: Roadmap ---
    add_content_slide("实施路线图 (Project Roadmap)", [
        "阶段一：MVP 原型 (1-2个月)",
        ("核心对话界面上线", 1),
        ("跑通 C++/Python 混合调用流程 (Benchmark)", 1),
        ("基础 RAG 问答功能", 1),
        "阶段二：平台化 (3-4个月)",
        ("可视化 Agent 编排界面", 1),
        ("多用户权限体系", 1),
        "阶段三：生态化 (6个月+)",
        ("开放 API，构建内部 Agent 技能市场", 1),
        ("整合第三方 ERP/CRM 系统", 1)
    ])

    # --- Slide 11: ROI ---
    add_content_slide("预期收益 (ROI)", [
        "效率提升：",
        ("复杂数据处理流程从“小时级”缩短至“分钟级”", 1),
        ("知识沉淀：", 1),
        ("将散落在员工电脑里的文档变为可检索的企业资产", 1),
        ("业务创新：", 1),
        ("快速生成各类分析报告、PPT，支持复杂决策辅助", 1)
    ])

    # --- Slide 12: End ---
    add_title_slide(
        "让 AI 成为企业的核心生产力",
        "Q & A\n\n感谢聆听"
    )

    # Save presentation
    output_filename = "Enterprise_Agent_Platform_Proposal.pptx"
    try:
        prs.save(output_filename)
        print(f"Successfully generated: {output_filename}")
    except PermissionError:
        print(f"Error: Could not save to {output_filename}. Please close the file if it is open.")

if __name__ == "__main__":
    create_presentation()
