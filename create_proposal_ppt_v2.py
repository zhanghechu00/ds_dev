from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

def create_presentation():
    # Colors
    DARK_BLUE = RGBColor(16, 42, 67)      # #102A43 (Deep Navy)
    TEAL = RGBColor(62, 166, 255)         # #3EA6FF (Bright Blue/Teal)
    LIGHT_GREY = RGBColor(240, 244, 248)  # #F0F4F8
    WHITE = RGBColor(255, 255, 255)
    TEXT_GREY = RGBColor(51, 51, 51)
    
    prs = Presentation()
    
    # Set slide dimensions to Widescreen (16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_geometric_accents(slide, is_title=False):
        """Adds geometric shapes to the slide for visual interest."""
        shapes = slide.shapes
        
        if is_title:
            # Full background
            bg = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
            bg.fill.solid()
            bg.fill.fore_color.rgb = DARK_BLUE
            bg.line.fill.background() # No line

            # Angle shape top right
            tri = shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(10), 0, Inches(3.33), Inches(3.33))
            tri.fill.solid()
            tri.fill.fore_color.rgb = TEAL
            tri.line.fill.background()
            tri.rotation = 0

            # Stripe bottom left
            rect = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.5), Inches(4), Inches(0.2))
            rect.fill.solid()
            rect.fill.fore_color.rgb = TEAL
            rect.line.fill.background()
            
        else:
            # Header Bar
            header = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
            header.fill.solid()
            header.fill.fore_color.rgb = DARK_BLUE
            header.line.fill.background()

            # Accent Line under header
            line = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.2), prs.slide_width, Inches(0.05))
            line.fill.solid()
            line.fill.fore_color.rgb = TEAL
            line.line.fill.background()

            # Small decorative square bottom right
            sq = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.8), Inches(6.8), Inches(0.3), Inches(0.3))
            sq.fill.solid()
            sq.fill.fore_color.rgb = TEAL
            sq.line.fill.background()

    def add_title_slide(title_text, subtitle_text):
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # 6 is Blank
        add_geometric_accents(slide, is_title=True)
        
        # Title
        left = Inches(1)
        top = Inches(2.5)
        width = Inches(11)
        height = Inches(2)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p = tf.add_paragraph()
        p.text = title_text
        p.font.size = Pt(54)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = 'Microsoft YaHei' # Or Arial
        
        # Subtitle
        top = Inches(4.5)
        txBox2 = slide.shapes.add_textbox(left, top, width, height)
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        
        for line in subtitle_text.split('\n'):
            p = tf2.add_paragraph()
            p.text = line
            p.font.size = Pt(24)
            p.font.color.rgb = LIGHT_GREY
            p.font.name = 'Microsoft YaHei'

    def add_content_slide(title_text, content_list, highlight_box=False):
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank
        add_geometric_accents(slide, is_title=False)
        
        # Title in Header
        left = Inches(0.5)
        top = Inches(0.2)
        width = Inches(12)
        height = Inches(1)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        
        p = tf.add_paragraph()
        p.text = title_text
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = 'Microsoft YaHei'
        p.alignment = PP_ALIGN.LEFT
        
        # Content Body
        left = Inches(0.8)
        top = Inches(1.8)
        width = Inches(11.5)
        height = Inches(5)
        
        # Optional: Add a subtle backing box for content
        if highlight_box:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left - Inches(0.2), top - Inches(0.2), width + Inches(0.4), height + Inches(0.2))
            box.fill.solid()
            box.fill.fore_color.rgb = LIGHT_GREY
            box.line.fill.background()
            # Move box to back? python-pptx z-order follows creation order. 
            # Since we add text after, text will be on top.
        
        txBoxBody = slide.shapes.add_textbox(left, top, width, height)
        tf = txBoxBody.text_frame
        tf.word_wrap = True
        
        for item in content_list:
            if isinstance(item, tuple):
                text = item[0]
                level = item[1]
            else:
                text = item
                level = 0
            
            p = tf.add_paragraph()
            p.text = text
            p.level = level
            p.font.name = 'Microsoft YaHei'
            p.font.size = Pt(28 - (level * 4))
            p.font.color.rgb = TEXT_GREY
            p.space_after = Pt(12)
            p.space_before = Pt(6)

    # --- Slide 1: Cover ---
    add_title_slide(
        "企业级 AI Agent 智能体平台建设方案",
        "构建“数字员工”生产力引擎，驱动业务智能化转型\n\n汇报人：[您的名字/团队名称]\n日期：2026年1月"
    )

    # --- Slide 2: Executive Summary ---
    add_content_slide("核心愿景 (Executive Summary)", [
        "目标：打造企业统一的 AI 智能体开发与运行平台",
        "核心价值主张：",
        ("低门槛：让业务人员也能“雇佣”数字员工", 1),
        ("破孤岛：打通文档、数据库与遗留系统", 1),
        ("高效率：实现从“对话”到“行动”的自动化闭环", 1),
        "平台关键产出：",
        ("Agent 开发工坊 (可视编排)", 1),
        ("工具连接网关 (MCP Gateway)", 1),
        ("知识中台 (Knowledge Hub)", 1)
    ])

    # --- Slide 3: Why Now ---
    add_content_slide("现状与痛点 (Why Now?)", [
        "行业趋势：大模型从 Chat (闲聊) 进化为 Agent (行动)",
        "当前企业痛点：",
        ("能力割裂：文档知识与业务系统无法互通", 1),
        ("  (例：LLM 无法理解 C++ 处理的油气数据)", 2),
        ("僵化：传统 RPA 无法处理模糊指令", 1),
        ("  (例：“分析这份报告，如果没有结论则忽略”)", 2),
        ("管控缺失：缺乏统一的安全审计与权限管理", 1)
    ])

    # --- Slide 4: Architecture ---
    add_content_slide("平台总体架构 (High-Level Architecture)", [
        "架构设计分层：",
        ("应用层：智能客服、自动化办公、专业数据分析助手", 1),
        ("功能层：Agent 开发工坊、应用市场、管理后台", 1),
        ("引擎层 (Brain)：ReAct 编排 + MCP 工具网关 + 记忆系统", 1),
        ("模型层 (MaaS)：GPT-4 / DeepSeek-70B (混合部署)", 1),
        "",
        "[请在此处插入架构图 platform_architecture.png]"
    ], highlight_box=True)

    # --- Slide 5: Tool Gateway (MCP) ---
    add_content_slide("核心引擎亮点：通用工具网关", [
        "痛点解决：如何让 AI 调用本地复杂的 Python/C++ 业务代码？",
        "技术方案：基于 MCP (Model Context Protocol) 协议",
        "案例演示：",
        ("AI 智能体 -> 编排引擎", 1),
        ("-> 直接调度 backend.py (Python)", 1),
        ("-> 调用 ExportPetroFileV32.exe (遗留 C++ 模块)", 1),
        ("优势：解耦大模型与底层业务逻辑，支持异构语言环境", 0)
    ], highlight_box=True)

    # --- Slide 6: Knowledge Engine (RAG) ---
    add_content_slide("知识中台：RAG 检索增强", [
        "功能：让模型拥有企业私有知识",
        "流程：文档上传 -> 自动切片 -> 向量化存储 -> 语义检索 -> 回答生成",
        "特色能力：",
        ("支持多格式解析 (PDF, MD, Word, 代码)", 0),
        ("源头溯源：回答中包含原文引用链接，确保可信度", 0),
        ("知识应用：如“基于技术文档自动生成 PPT”", 0)
    ])

    # --- Slide 7: Tech Stack ---
    add_content_slide("技术栈选型 (Tech Stack)", [
        "后端：Python (FastAPI) + LangChain / LangGraph",
        "前端：React + 流程编排组件",
        "向量数据库：PGVector (推荐) 或 Milvus",
        "核心协议：",
        ("全面适配 MCP (Model Context Protocol)", 1),
        ("确保工具生态兼容性，方便未来扩展", 1),
        "部署架构：Docker 容器化 + K8s 集群"
    ])

    # --- Slide 8: Hardware Resources ---
    add_content_slide("基础设施配置 (Resource Requirements)", [
        "方案 A：混合云/API 模式 (推荐 - 高性价比)",
        ("应用服务器 (2台)：运行 Agent 逻辑、MCP Server", 1),
        ("向量库服务器 (1台)：大内存，存储企业知识库", 1),
        ("遗留系统节点 (1台 Windows)：运行旧版 C++ 工具", 1),
        "方案 B：私有化部署 (数据严控)",
        ("LLM 推理集群：需配备 NVIDIA A100/H800 或 RTX 4090 集群", 1),
        ("显存需求：支持 70B+ 参数模型 (如 DeepSeek-70B)", 1)
    ])

    # --- Slide 9: Budget ---
    add_content_slide("成本构成与预算 (Budget Breakdown)", [
        "一次性投入 (CAPEX)：",
        ("硬件：服务器与 GPU 采购费用", 1),
        ("软件：OS、数据库授权、数据清洗人力", 1),
        ("研发人力：架构师、前后端、Prompt 工程师", 1),
        "持续运营 (OPEX)：",
        ("Token 费用：公有云 API 调用费 (按量)", 1),
        ("能耗与维保：私有化机房电费、系统运维", 1),
        ("持续迭代：Agent 技能更新、知识库维护", 1)
    ])

    # --- Slide 10: Roadmap ---
    add_content_slide("实施路线图 (Project Roadmap)", [
        "阶段一：MVP 原型 (1-2个月)",
        ("核心对话界面上线", 1),
        ("跑通 C++/Python 混合调用流程 (Benchmark)", 1),
        ("基础 RAG 问答功能", 1),
        "阶段二：平台化 (3-4个月)",
        ("可视化 Agent 编排界面", 1),
        ("多用户权限体系", 1),
        "阶段三：生态化 (6个月+)",
        ("开放 API，构建内部 Agent 技能市场", 1),
        ("整合第三方 ERP/CRM 系统", 1)
    ])

    # --- Slide 11: ROI ---
    add_content_slide("预期收益 (ROI)", [
        "效率提升：",
        ("复杂数据处理流程从“小时级”缩短至“分钟级”", 1),
        ("知识沉淀：", 1),
        ("将散落在员工电脑里的文档变为可检索的企业资产", 1),
        ("业务创新：", 1),
        ("快速生成各类分析报告、PPT，支持复杂决策辅助", 1)
    ])

    # --- Slide 12: End ---
    add_title_slide(
        "让 AI 成为企业的核心生产力",
        "Q & A\n\n感谢聆听"
    )

    # Save presentation
    output_filename = "Enterprise_Agent_Platform_Proposal_v2.pptx"
    try:
        prs.save(output_filename)
        print(f"Successfully generated: {output_filename}")
    except PermissionError:
        print(f"Error: Could not save to {output_filename}. Please close the file if it is open.")

if __name__ == "__main__":
    create_presentation()
