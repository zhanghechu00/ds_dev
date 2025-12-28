from docx import Document

def create_test_doc():
    doc = Document()
    doc.add_heading('地质参数一致性测试报告', 0)

    # Case 1: Consistent (一致)
    doc.add_heading('1. 孔隙度测试 (一致)', level=1)
    doc.add_paragraph('根据实验室测定，该岩心的孔隙度为：15.5%。')
    
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '参数名称'
    hdr_cells[1].text = '测定值'
    
    row_cells = table.rows[1].cells
    row_cells[0].text = '孔隙度'
    row_cells[1].text = '15.5%'

    doc.add_paragraph('\n')

    # Case 2: Inconsistent (不一致)
    doc.add_heading('2. 渗透率测试 (不一致)', level=1)
    doc.add_paragraph('报告正文中记录的渗透率为：25.0mD。')

    table2 = doc.add_table(rows=2, cols=2)
    table2.style = 'Table Grid'
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = '参数名称'
    hdr_cells2[1].text = '测定值'
    
    row_cells2 = table2.rows[1].cells
    row_cells2[0].text = '渗透率'
    row_cells2[1].text = '20.0mD'

    doc.save('test_consistency.docx')
    print("test_consistency.docx created successfully.")

if __name__ == "__main__":
    create_test_doc()
