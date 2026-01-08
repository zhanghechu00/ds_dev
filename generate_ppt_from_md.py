import argparse
import os
import re
from dataclasses import dataclass, field
from typing import List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@dataclass
class MdSection:
    title: str
    intro_lines: List[str] = field(default_factory=list)
    subsections: List[str] = field(default_factory=list)
    bullet_lines: List[str] = field(default_factory=list)
    code_blocks: List[str] = field(default_factory=list)
    tables: List[List[List[str]]] = field(default_factory=list)  # list of tables -> rows -> cells


def _clean_text(line: str) -> str:
    return re.sub(r"\s+", " ", line.strip())


def _looks_like_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def _is_table_separator_row(line: str) -> bool:
    s = line.strip()
    if not _looks_like_table_row(s):
        return False
    # e.g. |---|---| or | --- | :---: |
    cells = [c.strip() for c in s.strip("|").split("|")]
    if not cells:
        return False
    for c in cells:
        if not c:
            return False
        if not re.fullmatch(r"[:\-\s]+", c):
            return False
    return True


def parse_markdown(md_text: str, fallback_title: str = "(Untitled)") -> tuple[str, List[MdSection]]:
    title = "(Untitled)"
    sections: List[MdSection] = []
    current: Optional[MdSection] = None

    preface_lines: List[str] = []

    in_code = False
    code_lines: List[str] = []
    code_lang: str = ""

    # table parsing state (within a section)
    pending_table: List[List[str]] = []
    table_header_seen = False

    for raw in md_text.splitlines():
        line = raw.rstrip("\n")
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
                code_lang = stripped[3:].strip()
            else:
                in_code = False
                if current is not None:
                    # Keep the full code block; language is optional
                    header = f"```{code_lang}".rstrip()
                    body = "\n".join(code_lines).rstrip()
                    current.code_blocks.append((header + "\n" + body + "\n```") if body else (header + "\n```"))
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not stripped or stripped == "---":
            continue

        # Support "bold headings" used as section titles, e.g. **1) ...** or **A. ...**
        bold_heading = re.match(r"^\*\*(.+)\*\*$", stripped)
        if bold_heading:
            heading_text = _clean_text(bold_heading.group(1))
            if heading_text:
                # Create a new section when seeing a bold heading
                current = MdSection(title=heading_text)
                sections.append(current)
                continue

        # Table parsing: capture consecutive table rows
        if current is not None and _looks_like_table_row(stripped):
            # detect separator to decide header presence
            if _is_table_separator_row(stripped):
                table_header_seen = True
                continue

            row = [c.strip() for c in stripped.strip("|").split("|")]
            pending_table.append(row)
            continue
        else:
            # flush pending table when leaving table region
            if current is not None and pending_table:
                # store one table
                current.tables.append(pending_table)
                pending_table = []
                table_header_seen = False

        if stripped.startswith("# "):
            title = _clean_text(stripped[2:]) or title
            continue

        if stripped.startswith("## "):
            current = MdSection(title=_clean_text(stripped[3:]))
            sections.append(current)
            continue

        if stripped.startswith("### "):
            if current is not None:
                current.subsections.append(_clean_text(stripped[4:]))
            continue

        if current is None:
            # Keep preface text (before first section) as an overview slide later
            cleaned_preface = _clean_text(stripped)
            if cleaned_preface:
                preface_lines.append(cleaned_preface)
            continue

        cleaned = _clean_text(stripped)

        # capture bullets / numbered list
        if re.match(r"^[-*]\s+", cleaned) or re.match(r"^\d+\.\s+", cleaned):
            current.bullet_lines.append(cleaned)
        else:
            current.intro_lines.append(cleaned)

    if title == "(Untitled)":
        title = fallback_title or title

    # If the doc has preface text, put it into an overview section at the beginning.
    if preface_lines:
        overview = MdSection(title="概览", intro_lines=preface_lines[:12])
        if sections:
            sections.insert(0, overview)
        else:
            sections.append(overview)

    return title, sections


def create_presentation(md_title: str, sections: List[MdSection], subtitle: str, output_path: str) -> None:
    prs = Presentation()

    # Color scheme (same as create_training_ppt.py)
    COLOR_PRIMARY = RGBColor(0, 51, 102)  # #003366
    COLOR_ACCENT = RGBColor(51, 102, 153)  # #336699
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT = RGBColor(64, 64, 64)  # #404040

    def add_title_slide(title_text: str, subtitle_text: str):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_PRIMARY

        title_shape = slide.shapes.add_textbox(Inches(1), Inches(2.3), Inches(8), Inches(1.6))
        tf = title_shape.text_frame
        tf.word_wrap = True
        p = tf.add_paragraph()
        p.text = title_text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(42)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.name = "Microsoft YaHei"

        subtitle_shape = slide.shapes.add_textbox(Inches(1), Inches(4.1), Inches(8), Inches(1.6))
        stf = subtitle_shape.text_frame
        stf.word_wrap = True
        p2 = stf.add_paragraph()
        p2.text = subtitle_text
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(20)
        p2.font.color.rgb = RGBColor(200, 200, 200)
        p2.font.name = "Microsoft YaHei"

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7), Inches(10), Inches(0.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_ACCENT
        bar.line.fill.background()

    def add_content_slide(title_text: str, content_items: List[str]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.2))
        header.fill.solid()
        header.fill.fore_color.rgb = COLOR_PRIMARY
        header.line.fill.background()

        title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
        p = title_shape.text_frame.add_paragraph()
        p.text = title_text
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.name = "Microsoft YaHei"

        content_shape = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
        tf = content_shape.text_frame
        tf.word_wrap = True

        for item in content_items:
            if item is None:
                continue
            text = item.rstrip()
            if not text:
                continue
            para = tf.add_paragraph()
            para.text = text
            para.font.size = Pt(18)
            para.font.color.rgb = COLOR_TEXT
            para.font.name = "Microsoft YaHei"
            para.space_after = Pt(10)

            if text.strip().startswith("-") or re.match(r"^\d+\.\s+", text.strip()):
                para.level = 1
            else:
                para.level = 0

    def add_code_slide(title_text: str, code_text: str):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.2))
        header.fill.solid()
        header.fill.fore_color.rgb = COLOR_PRIMARY
        header.line.fill.background()

        title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
        p = title_shape.text_frame.add_paragraph()
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.name = "Microsoft YaHei"

        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(8.8), Inches(5.7))
        tf = box.text_frame
        tf.word_wrap = False
        tf.clear()
        p2 = tf.paragraphs[0]
        p2.text = code_text
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT
        p2.font.name = "Consolas"

    def add_table_slide(title_text: str, table_rows: List[List[str]]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.2))
        header.fill.solid()
        header.fill.fore_color.rgb = COLOR_PRIMARY
        header.line.fill.background()

        title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
        p = title_shape.text_frame.add_paragraph()
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.font.name = "Microsoft YaHei"

        if not table_rows:
            add_content_slide(title_text, ["（表格为空）"])
            return

        cols = max(len(r) for r in table_rows)
        rows = len(table_rows)

        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(5.3)

        table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
        table = table_shape.table

        # Basic styling
        for r in range(rows):
            for c in range(cols):
                cell = table.cell(r, c)
                cell.text = table_rows[r][c] if c < len(table_rows[r]) else ""
                for para in cell.text_frame.paragraphs:
                    para.font.name = "Microsoft YaHei"
                    para.font.size = Pt(14)
                    para.font.color.rgb = COLOR_TEXT
                # header row styling
                if r == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = COLOR_ACCENT
                    for para in cell.text_frame.paragraphs:
                        para.font.bold = True
                        para.font.color.rgb = COLOR_WHITE

    def chunk(items: List[str], max_items: int) -> List[List[str]]:
        return [items[i : i + max_items] for i in range(0, len(items), max_items)]

    add_title_slide(md_title, subtitle)

    for sec in sections:
        content: List[str] = []

        # Prefer a short intro (first 3 lines)
        if sec.intro_lines:
            content.extend(sec.intro_lines[:3])

        if sec.subsections:
            content.append("子主题：")
            content.extend([f"- {s}" for s in sec.subsections[:20]])
            if len(sec.subsections) > 20:
                content.append(f"- ...（共 {len(sec.subsections)} 个子主题）")

        # Include some bullets if present
        if sec.bullet_lines:
            content.append("要点：")
            content.extend(sec.bullet_lines[:16])
            if len(sec.bullet_lines) > 16:
                content.append(f"- ...（更多要点见文档，已省略）")

        if sec.code_blocks:
            content.append(f"代码示例：本节包含 {len(sec.code_blocks)} 处代码块（后续页展示）")

        if sec.tables:
            content.append(f"表格：本节包含 {len(sec.tables)} 个表格（后续页展示）")

        # Ensure there is at least something
        if not content:
            content = ["（本节主要为代码示例，PPT 未展开全文）"]

        # Split into multiple slides if too long
        pages = chunk(content, max_items=14)
        for idx, page in enumerate(pages, start=1):
            slide_title = sec.title if len(pages) == 1 else f"{sec.title}（{idx}/{len(pages)}）"
            add_content_slide(slide_title, page)

        # Add table slides (make summary crystal clear)
        if sec.tables:
            for t_idx, table_rows in enumerate(sec.tables, start=1):
                t_title = f"{sec.title}：表格" if len(sec.tables) == 1 else f"{sec.title}：表格 {t_idx}/{len(sec.tables)}"
                add_table_slide(t_title, table_rows)

        # Add code slides (more complete)
        if sec.code_blocks:
            for c_idx, code_text in enumerate(sec.code_blocks, start=1):
                c_title = f"{sec.title}：代码示例" if len(sec.code_blocks) == 1 else f"{sec.title}：代码示例 {c_idx}/{len(sec.code_blocks)}"
                # Avoid extremely long single textboxes; truncate very long blocks
                lines = code_text.splitlines()
                if len(lines) > 60:
                    code_text = "\n".join(lines[:60]) + "\n...（已截断，全文见 Markdown）"
                add_code_slide(c_title, code_text)

    prs.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate PPTX from a Markdown file using project slide style.")
    parser.add_argument("--md", required=True, help="Path to the Markdown file")
    parser.add_argument("--out", default="", help="Output .pptx path")
    args = parser.parse_args()

    md_path = os.path.abspath(args.md)
    if not os.path.exists(md_path):
        raise FileNotFoundError(f"Markdown file not found: {md_path}")

    with open(md_path, "r", encoding="utf-8", errors="ignore") as f:
        md_text = f.read()

    fallback_title = os.path.splitext(os.path.basename(md_path))[0]
    md_title, sections = parse_markdown(md_text, fallback_title=fallback_title)

    out_path = args.out.strip() or os.path.join(
        os.path.dirname(md_path),
        f"{os.path.splitext(os.path.basename(md_path))[0]}_AUTO.pptx",
    )

    subtitle = "基于项目脚本风格自动生成\n（内容来自 Markdown 文档）"
    create_presentation(md_title, sections, subtitle, out_path)
    print(f"PPT generated: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
